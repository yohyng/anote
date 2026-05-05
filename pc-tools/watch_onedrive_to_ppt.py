"""
Watch OneDrive/PencilRoom/inbox and append new PNG files to an existing PPTX.

Install:
    pip install python-pptx watchdog

Example:
    python watch_onedrive_to_ppt.py ^
      --inbox "C:\Users\YOURNAME\OneDrive\PencilRoom\inbox" ^
      --pptx "C:\Users\YOURNAME\OneDrive\PencilRoom\deck.pptx" ^
      --processed "C:\Users\YOURNAME\OneDrive\PencilRoom\processed"
"""

from __future__ import annotations

import argparse
import shutil
import time
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches
from watchdog.events import FileSystemEventHandler
from watchdog.observers import Observer


PNG_EXTENSIONS = {".png"}


def wait_until_file_stable(path: Path, checks: int = 5, interval: float = 0.5) -> bool:
    previous_size = -1
    stable_count = 0

    for _ in range(checks * 4):
        if not path.exists():
            time.sleep(interval)
            continue

        size = path.stat().st_size
        if size == previous_size and size > 0:
            stable_count += 1
        else:
            stable_count = 0

        if stable_count >= checks:
            return True

        previous_size = size
        time.sleep(interval)

    return False


def append_png_to_pptx(pptx_path: Path, png_path: Path) -> None:
    if pptx_path.exists():
        prs = Presentation(str(pptx_path))
    else:
        prs = Presentation()
        prs.slide_width = Inches(13.333333)
        prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    slide_w = prs.slide_width
    slide_h = prs.slide_height

    slide.shapes.add_picture(str(png_path), 0, 0, width=slide_w, height=slide_h)
    prs.save(str(pptx_path))


def process_png(path: Path, pptx_path: Path, processed_dir: Path) -> None:
    if path.suffix.lower() not in PNG_EXTENSIONS:
        return

    if not wait_until_file_stable(path):
        print(f"[skip] file not stable: {path}")
        return

    print(f"[append] {path.name}")
    append_png_to_pptx(pptx_path, path)

    processed_dir.mkdir(parents=True, exist_ok=True)
    destination = processed_dir / path.name
    if destination.exists():
        destination = processed_dir / f"{path.stem}_{int(time.time())}{path.suffix}"

    shutil.move(str(path), str(destination))
    print(f"[done] moved to {destination}")


class InboxHandler(FileSystemEventHandler):
    def __init__(self, pptx_path: Path, processed_dir: Path):
        super().__init__()
        self.pptx_path = pptx_path
        self.processed_dir = processed_dir

    def on_created(self, event):
        if event.is_directory:
            return
        process_png(Path(event.src_path), self.pptx_path, self.processed_dir)

    def on_moved(self, event):
        if event.is_directory:
            return
        process_png(Path(event.dest_path), self.pptx_path, self.processed_dir)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--inbox", required=True, help="OneDrive-synced inbox folder containing PNG files.")
    parser.add_argument("--pptx", required=True, help="Target PPTX path.")
    parser.add_argument("--processed", required=True, help="Folder to move processed PNG files into.")
    parser.add_argument("--process-existing", action="store_true", help="Process existing PNG files before watching.")
    args = parser.parse_args()

    inbox_dir = Path(args.inbox)
    pptx_path = Path(args.pptx)
    processed_dir = Path(args.processed)

    inbox_dir.mkdir(parents=True, exist_ok=True)
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    processed_dir.mkdir(parents=True, exist_ok=True)

    if args.process_existing:
        for png_path in sorted(inbox_dir.glob("*.png")):
            process_png(png_path, pptx_path, processed_dir)

    observer = Observer()
    observer.schedule(InboxHandler(pptx_path, processed_dir), str(inbox_dir), recursive=False)
    observer.start()

    print(f"[watch] inbox: {inbox_dir}")
    print(f"[pptx]  target: {pptx_path}")
    print("Press Ctrl+C to stop.")

    try:
        while True:
            time.sleep(1)
    except KeyboardInterrupt:
        observer.stop()

    observer.join()


if __name__ == "__main__":
    main()
